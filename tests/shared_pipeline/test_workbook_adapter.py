import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from apc.config import Settings
from apc.db import LedgerDB
from apc.stages.container_expansion import run_container_expansion


STAMP = "2026-10-27T12:00:00Z"


def make_workbooks(folder):
    folder.mkdir(parents=True, exist_ok=True)

    valid = folder / "Workbook_Valid.xlsx"
    corrupt = folder / "Workbook_Corrupt.xlsx"

    book = Workbook()
    visible = book.active
    visible.title = "Visible"
    visible.append(["Name", "Value"])
    visible.append(["Alpha", "100"])

    hidden = book.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Hidden", "200"])

    very_hidden = book.create_sheet("VeryHidden")
    very_hidden.sheet_state = "veryHidden"
    very_hidden.append(["VeryHidden", "300"])

    book.create_sheet("Blank")
    book.save(valid)

    corrupt.write_bytes(b"NOT_A_VALID_EXCEL_WORKBOOK")
    return valid, corrupt


def insert_file(db, job_id, path, file_id):
    db.execute(
        """
        INSERT INTO file_processing_metrics (
            file_id, matter_id, job_id,
            original_path, normalized_path,
            extension, mime_type, source_bytes,
            is_container, container_depth,
            stage_status_json, exception_json,
            created_at, updated_at
        )
        VALUES (
            ?, 'MATTER-TEST', ?, ?, ?,
            'xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            ?, 0, 0, '{}', '[]', ?, ?
        )
        """,
        (
            file_id, job_id, str(path), path.name,
            path.stat().st_size, STAMP, STAMP,
        ),
    )


def run_case(base, enabled):
    folder = base / ("adapter" if enabled else "legacy")
    folder.mkdir(parents=True, exist_ok=True)
    inputs = folder / "input"

    valid, corrupt = make_workbooks(inputs)
    db = LedgerDB(str(folder / "ledger.sqlite"))
    db.init_schema()

    job_id = "JOB-ADAPTER" if enabled else "JOB-LEGACY"

    try:
        db.execute(
            """
            INSERT INTO processing_job (
                job_id, matter_id, client_id,
                created_at, status
            )
            VALUES (?, 'MATTER-TEST', 'CLIENT-TEST', ?, 'pending')
            """,
            (job_id, STAMP),
        )

        insert_file(db, job_id, valid, "FILE-GOOD")
        insert_file(db, job_id, corrupt, "FILE-BAD")

        with patch.dict(
            os.environ,
            {
                "INSYT_WORKBOOK_ADAPTER_ENABLED":
                    "1" if enabled else "0"
            },
        ):
            run_container_expansion(
                db=db,
                settings=Settings(),
                job_id=job_id,
                matter_id="MATTER-TEST",
                input_dir=str(inputs),
            )

        rows = db.query(
            """
            SELECT *
            FROM file_processing_metrics
            WHERE job_id=?
            ORDER BY normalized_path, file_id
            """,
            (job_id,),
        )

        children = {}
        parents = {}

        for row in rows:
            stage = json.loads(row["stage_status_json"] or "{}")

            if row["parent_file_id"] == "FILE-GOOD":
                content_path = Path(row["original_path"])
                assert content_path.is_file()

                children[row["normalized_path"]] = {
                    "content": content_path.read_bytes(),
                    "extension": row["extension"],
                    "file_id": row["file_id"],
                    "parent": row["parent_file_id"],
                    "lineage": stage.get("workbook_sheet"),
                }

            if row["file_id"] in {"FILE-GOOD", "FILE-BAD"}:
                parents[row["file_id"]] = {
                    "container": row["is_container"],
                    "stage": stage,
                    "exceptions": json.loads(
                        row["exception_json"] or "[]"
                    ),
                }

        assert valid.is_file()
        assert corrupt.read_bytes() == b"NOT_A_VALID_EXCEL_WORKBOOK"

        return children, parents
    finally:
        db.close()


class WorkbookAdapterRegression(unittest.TestCase):

    def test_legacy_registered_ledger_parity(self):
        with tempfile.TemporaryDirectory(
            prefix="insyt_workbook_regression_"
        ) as temp:
            root = Path(temp)
            old_children, old_parents = run_case(root, False)
            new_children, new_parents = run_case(root, True)

            self.assertEqual(len(old_children), 3)
            self.assertEqual(len(new_children), 3)
            self.assertEqual(
                set(old_children),
                set(new_children),
            )

            ids = [c["file_id"] for c in new_children.values()]
            self.assertEqual(len(ids), len(set(ids)))

            for path, old in old_children.items():
                new = new_children[path]

                self.assertEqual(old["content"], new["content"])
                self.assertEqual(old["extension"], "csv")
                self.assertEqual(new["extension"], "csv")
                self.assertEqual(new["parent"], "FILE-GOOD")
                self.assertIsNotNone(new["lineage"])
                self.assertEqual(
                    new["lineage"]["original_workbook_file_id"],
                    "FILE-GOOD",
                )

                for field in (
                    "sheet_name", "sheet_index", "sheet_visibility",
                    "sheet_rows_seen", "sheet_nonblank_row_count",
                    "sheet_column_count", "sheet_nonblank_cell_count",
                ):
                    self.assertEqual(
                        old["lineage"].get(field),
                        new["lineage"].get(field),
                        f"{path}: {field}",
                    )

            old_manifest = old_parents["FILE-GOOD"][
                "stage"
            ]["workbook_expansion"]

            new_manifest = new_parents["FILE-GOOD"][
                "stage"
            ]["workbook_expansion"]

            for field in (
                "sheet_count", "nonblank_sheet_count",
                "blank_sheet_count", "generated_child_count",
            ):
                self.assertEqual(
                    old_manifest[field], new_manifest[field], field
                )

            def sheets(manifest):
                return [
                    {
                        k: value.get(k)
                        for k in (
                            "sheet_index", "sheet_name", "visibility",
                            "status", "rows_seen",
                            "nonblank_row_count",
                            "max_column_count",
                            "nonblank_cell_count",
                        )
                    }
                    for value in manifest["sheets"]
                ]

            self.assertEqual(
                sheets(old_manifest), sheets(new_manifest)
            )

            for parents in (old_parents, new_parents):
                self.assertEqual(
                    parents["FILE-GOOD"]["container"], 1
                )
                self.assertEqual(
                    parents["FILE-BAD"]["container"], 0
                )
                self.assertTrue(
                    parents["FILE-BAD"]["exceptions"]
                )


if __name__ == "__main__":
    unittest.main()
