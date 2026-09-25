from __future__ import annotations

import os

from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .rules import DetectionRule
from ..azure_blob_adapter import DualStorageBlobAdapter
from ..azure_layout import AzureRoutingConfig


DEFAULT_CAPTURE_RULE_SHEETS = (
    "General PII",
)


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _normalize_header(value: Any) -> str:
    return (
        _clean_text(value)
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
        .replace("/", "_")
    )


def _parse_bool(
    value: Any,
    *,
    default: bool = True,
) -> bool:
    if value is None:
        return default

    if isinstance(value, bool):
        return value

    text = _clean_text(value).lower()

    if not text:
        return default

    if text in {
        "1",
        "true",
        "yes",
        "y",
        "enabled",
        "enable",
        "active",
    }:
        return True

    if text in {
        "0",
        "false",
        "no",
        "n",
        "disabled",
        "disable",
        "inactive",
    }:
        return False

    return default


def _parse_float(
    value: Any,
    *,
    default: float = 0.70,
) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return default

    return max(
        0.0,
        min(
            1.0,
            result,
        ),
    )


def _parse_list(
    value: Any,
) -> tuple[str, ...]:
    text = _clean_text(value)

    if not text:
        return ()

    normalized = (
        text
        .replace("\r\n", "|")
        .replace("\n", "|")
        .replace(";", "|")
    )

    return tuple(
        item.strip()
        for item in normalized.split("|")
        if item.strip()
    )


def _row_value(
    row: dict[str, Any],
    *names: str,
) -> Any:
    for name in names:
        normalized = _normalize_header(
            name
        )

        if normalized in row:
            value = row.get(
                normalized
            )

            if value is not None:
                return value

    return None


def _rule_from_row(
    row: dict[str, Any],
    *,
    row_number: int,
) -> DetectionRule | None:
    rule_id = _clean_text(
        _row_value(
            row,
            "rule_id",
            "rule id",
            "id",
        )
    )

    entity_type = _clean_text(
        _row_value(
            row,
            "entity_type",
            "entity type",
            "criterion",
            "capture criterion",
            "category",
        )
    )

    regex_pattern = _clean_text(
        _row_value(
            row,
            "regex_pattern",
            "regex pattern",
            "regex",
            "pattern",
        )
    )

    enabled = _parse_bool(
        _row_value(
            row,
            "enabled",
            "active",
        ),
        default=True,
    )

    #
    # Completely blank workbook rows are ignored.
    #
    if (
        not rule_id
        and not entity_type
        and not regex_pattern
    ):
        return None

    if not rule_id:
        rule_id = (
            f"CAPTURE-{row_number:04d}"
        )

    if not entity_type:
        raise ValueError(
            f"Capture Detection rule {rule_id} "
            "is missing Entity Type / "
            "Capture Criterion."
        )

    if enabled and not regex_pattern:
        raise ValueError(
            f"Capture Detection rule {rule_id} "
            "is enabled but has no Regex Pattern."
        )

    entity_subtype = _clean_text(
        _row_value(
            row,
            "entity_subtype",
            "entity subtype",
            "subtype",
            "issue",
        )
    )

    context_terms = _parse_list(
        _row_value(
            row,
            "context_terms",
            "context terms",
            "context",
            "keywords",
        )
    )

    validator = _clean_text(
        _row_value(
            row,
            "validator",
            "validation",
            "validation method",
        )
    )

    base_confidence = _parse_float(
        _row_value(
            row,
            "base_confidence",
            "base confidence",
            "confidence",
            "weight",
        ),
        default=0.70,
    )

    framework = _parse_list(
        _row_value(
            row,
            "framework",
            "frameworks",
            "issue group",
            "issue groups",
        )
    )

    methods = _parse_list(
        _row_value(
            row,
            "methods",
            "method",
            "match type",
        )
    )

    if not methods:
        methods = (
            "regex",
        )

    metadata: dict[str, Any] = {
        "workspace": "capture",
        "source": (
            "Capture_Search_Rules.xlsx"
        ),
        "workbook_row": row_number,
    }

    capture_group_value = _row_value(
        row,
        "capture_group",
        "capture group",
        "regex capture group",
    )

    if capture_group_value not in (
        None,
        "",
    ):
        try:
            capture_group = int(
                float(
                    str(
                        capture_group_value
                    ).strip()
                )
            )
        except (
            TypeError,
            ValueError,
        ):
            raise ValueError(
                "Invalid Capture Group "
                f"at workbook row "
                f"{row_number}: "
                f"{capture_group_value!r}"
            )

        if capture_group < 0:
            raise ValueError(
                "Capture Group cannot be "
                "negative at workbook row "
                f"{row_number}."
            )

        metadata[
            "capture_group"
        ] = capture_group

    result_coding = _clean_text(
        _row_value(
            row,
            "coding_result",
            "coding result",
            "result_coding",
            "result coding",
        )
    )

    if result_coding:
        metadata[
            "coding_result"
        ] = result_coding

    action = _clean_text(
        _row_value(
            row,
            "action",
            "include_exclude",
            "include exclude",
            "disposition",
        )
    )

    if action:
        metadata[
            "action"
        ] = action

    threshold = _row_value(
        row,
        "threshold",
        "responsive_threshold",
        "responsive threshold",
    )

    if threshold not in (
        None,
        "",
    ):
        metadata[
            "threshold"
        ] = _parse_float(
            threshold,
            default=base_confidence,
        )

    return DetectionRule(
        rule_id=rule_id,
        entity_type=entity_type,
        entity_subtype=entity_subtype,
        regex_pattern=regex_pattern,
        context_terms=context_terms,
        validator=validator,
        base_confidence=base_confidence,
        framework=framework,
        enabled=enabled,
        methods=methods,
        metadata=metadata,
    )


def _rules_from_sheet(
    worksheet,
) -> list[DetectionRule]:
    rows = worksheet.iter_rows(
        values_only=True
    )

    try:
        raw_headers = next(
            rows
        )
    except StopIteration:
        return []

    headers = [
        _normalize_header(
            value
        )
        for value in raw_headers
    ]

    if not any(headers):
        return []

    rules: list[
        DetectionRule
    ] = []

    for row_number, values in enumerate(
        rows,
        start=2,
    ):
        row = {
            header: value
            for header, value in zip(
                headers,
                values,
            )
            if header
        }

        rule = _rule_from_row(
            row,
            row_number=row_number,
        )

        if rule is not None:
            rules.append(
                rule
            )

    return rules


def _rules_from_workbook(
    workbook,
    *,
    sheet_names: list[str] | tuple[str, ...] | None = None,
) -> list[DetectionRule]:
    requested_sheets = [
        _clean_text(
            sheet_name
        )
        for sheet_name in (
            sheet_names
            or DEFAULT_CAPTURE_RULE_SHEETS
        )
        if _clean_text(
            sheet_name
        )
    ]

    if not requested_sheets:
        requested_sheets = list(
            DEFAULT_CAPTURE_RULE_SHEETS
        )

    missing_sheets = [
        sheet_name
        for sheet_name in requested_sheets
        if sheet_name not in workbook.sheetnames
    ]

    if missing_sheets:
        raise ValueError(
            "Capture Search Rules workbook is missing "
            "requested sheet(s): "
            + ", ".join(
                missing_sheets
            )
        )

    rules_by_id: dict[
        str,
        DetectionRule
    ] = {}

    for sheet_name in requested_sheets:
        worksheet = workbook[
            sheet_name
        ]

        sheet_rules = (
            _rules_from_sheet(
                worksheet
            )
        )

        for rule in sheet_rules:
            rule_id = str(
                rule.rule_id
                or ""
            ).strip()

            if not rule_id:
                continue

            if rule_id in rules_by_id:
                raise ValueError(
                    "Duplicate Capture Search Rule ID "
                    f"{rule_id!r} found while loading "
                    f"sheet {sheet_name!r}."
                )

            rules_by_id[
                rule_id
            ] = rule

    return list(
        rules_by_id.values()
    )


def load_capture_rules_from_bytes(
    workbook_bytes: bytes,
    *,
    sheet_names: list[str] | tuple[str, ...] | None = None,
) -> list[DetectionRule]:
    if not workbook_bytes:
        return []

    workbook = load_workbook(
        filename=BytesIO(
            workbook_bytes
        ),
        read_only=True,
        data_only=True,
    )

    try:
        return _rules_from_workbook(
            workbook,
            sheet_names=sheet_names,
        )
    finally:
        workbook.close()


def load_capture_rules_from_file(
    path: str | Path,
    *,
    sheet_names: list[str] | tuple[str, ...] | None = None,
) -> list[DetectionRule]:
    workbook_path = Path(
        path
    )

    if not workbook_path.exists():
        raise FileNotFoundError(
            workbook_path
        )

    return (
        load_capture_rules_from_bytes(
            workbook_path.read_bytes(),
            sheet_names=sheet_names,
        )
    )

def capture_rule_blob_candidates(
    *,
    client: str,
    project: str,
) -> list[str]:
    """
    Ordered Capture Detection rule workbook locations.

    Project-specific rules take precedence over system defaults.
    """

    routing = AzureRoutingConfig.from_args(
        workspace="capture",
        client=client,
        project=project,
        review_container=(
            os.getenv(
                "INSYT_REVIEW_CONTAINER_CAPTURE"
            )
            or None
        ),
    )

    project_prefix = routing.prefix
    project_key = routing.project

    return [
        (
            f"{project_prefix}/source/protocol/"
            f"{project_key}_Capture_Search_Rules.xlsx"
        ),
        (
            f"{project_prefix}/source/protocol/"
            "Capture_Search_Rules.xlsx"
        ),
        (
            "_system/protocol_templates/capture/"
            "Capture_Search_Rules.xlsx"
        ),
        (
            "System/ProtocolTemplates/"
            "Capture_Search_Rules.xlsx"
        ),
        (
            "system/ProtocolTemplates/"
            "Capture_Search_Rules.xlsx"
        ),
    ]


def load_capture_rules_from_azure(
    *,
    client: str,
    project: str,
    sheet_names: list[str] | tuple[str, ...] | None = None,
) -> tuple[list[DetectionRule], str]:
    """
    Load the first available Capture Detection workbook.

    Returns:
        (rules, blob_path)

    Project-specific workbooks override system defaults.
    """

    clean_client = str(
        client or ""
    ).strip()

    clean_project = str(
        project or ""
    ).strip()

    if not clean_client:
        raise ValueError(
            "Capture Detection rules require client."
        )

    if not clean_project:
        raise ValueError(
            "Capture Detection rules require project."
        )

    routing = AzureRoutingConfig.from_args(
        workspace="capture",
        client=clean_client,
        project=clean_project,
        review_container=(
            os.getenv(
                "INSYT_REVIEW_CONTAINER_CAPTURE"
            )
            or None
        ),
    )

    adapter = DualStorageBlobAdapter(
        routing
    )

    attempted_paths: list[str] = []

    for blob_path in (
        capture_rule_blob_candidates(
            client=clean_client,
            project=clean_project,
        )
    ):
        attempted_paths.append(
            blob_path
        )

        blob_client = (
            adapter.review_container
            .get_blob_client(
                blob_path
            )
        )

        if not blob_client.exists():
            continue

        workbook_bytes = (
            blob_client
            .download_blob()
            .readall()
        )

        rules = (
            load_capture_rules_from_bytes(
                workbook_bytes,
                sheet_names=sheet_names,
            )
        )

        if not rules:
            raise RuntimeError(
                "Capture Detection rule workbook "
                f"contains no usable rules: {blob_path}"
            )

        return (
            rules,
            blob_path,
        )

    raise FileNotFoundError(
        "Capture Detection rule workbook "
        "was not found. Checked: "
        + ", ".join(
            attempted_paths
        )
    )
