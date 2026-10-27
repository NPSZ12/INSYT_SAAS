from __future__ import annotations

import csv
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

OPENPYXL_EXTENSIONS = {"xlsx", "xlsm", "xltx", "xltm"}
XLRD_EXTENSIONS = {"xls"}
PYXLSB_EXTENSIONS = {"xlsb"}


def _normalize_extension(extension: str | None) -> str:
    return str(extension or "").strip().lower().lstrip(".")


def _safe_sheet_filename(
    sheet_name: str,
    sheet_index: int,
) -> str:
    clean = str(
        sheet_name
        or "Sheet"
    ).strip()

    clean = re.sub(
        r'[<>:"/\\|?*\x00-\x1f]',
        "_",
        clean,
    )

    clean = re.sub(
        r"\s+",
        "_",
        clean,
    )

    clean = clean.strip(
        " ._"
    )

    if not clean:
        clean = "Sheet"

    #
    # Keep generated path lengths sane.
    #
    clean = clean[:120]

    return (
        f"Sheet_{sheet_index:04d}"
        f"__{clean}.csv"
    )


def _cell_to_csv_value(
    value: Any,
) -> Any:
    """
    Normalize spreadsheet cell values for deterministic
    worksheet-derived CSV output.
    """

    if value is None:
        return ""

    if isinstance(
        value,
        datetime,
    ):
        return value.isoformat(
            sep=" "
        )

    if isinstance(
        value,
        date,
    ):
        return value.isoformat()

    if isinstance(
        value,
        time,
    ):
        return value.isoformat()

    if isinstance(
        value,
        bool,
    ):
        return (
            "TRUE"
            if value
            else "FALSE"
        )

    return value


def _trim_trailing_empty_cells(
    row_values: Iterable[Any],
) -> list[Any]:
    values = [
        _cell_to_csv_value(
            value
        )
        for value in row_values
    ]

    last_nonempty = -1

    for index, value in enumerate(
        values
    ):
        if str(
            value
        ).strip():
            last_nonempty = index

    if last_nonempty < 0:
        return []

    return values[
        : last_nonempty + 1
    ]


def _write_sheet_rows_to_csv(
    rows: Iterable[Iterable[Any]],
    csv_path: Path,
) -> dict[str, Any]:
    """
    Stream a worksheet to CSV.

    Rows that are completely blank are ignored.

    Trailing blank columns are omitted so heavily
    formatted spreadsheets do not generate enormous
    empty CSVs.
    """

    csv_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    row_count = 0
    nonblank_row_count = 0
    max_column_count = 0
    cell_value_count = 0

    #
    # utf-8-sig makes the derivative CSV friendlier to
    # Excel/Cyber Utility Suite while remaining valid UTF-8.
    #
    with csv_path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:
        writer = csv.writer(
            handle,
            lineterminator="\n",
        )

        for raw_row in rows:
            row_count += 1

            values = (
                _trim_trailing_empty_cells(
                    raw_row
                )
            )

            if not values:
                continue

            writer.writerow(
                values
            )

            nonblank_row_count += 1

            max_column_count = max(
                max_column_count,
                len(values),
            )

            cell_value_count += sum(
                1
                for value in values
                if str(
                    value
                ).strip()
            )

    is_blank = (
        nonblank_row_count == 0
    )

    if is_blank:
        try:
            csv_path.unlink()
        except FileNotFoundError:
            pass

        size_bytes = 0

    else:
        size_bytes = (
            csv_path
            .stat()
            .st_size
        )

    return {
        "is_blank": is_blank,
        "rows_seen": row_count,
        "nonblank_row_count": (
            nonblank_row_count
        ),
        "max_column_count": (
            max_column_count
        ),
        "nonblank_cell_count": (
            cell_value_count
        ),
        "csv_bytes": size_bytes,
    }


def _openpyxl_visibility(
    worksheet,
) -> str:
    value = str(
        getattr(
            worksheet,
            "sheet_state",
            "visible",
        )
        or "visible"
    ).strip().lower()

    if value in {
        "visible",
        "hidden",
        "veryhidden",
        "very_hidden",
    }:
        if value == "veryhidden":
            return "very_hidden"

        return value

    return "unknown"


def _xls_visibility(
    workbook,
    sheet_index_zero_based: int,
) -> str:
    #
    # xlrd may expose sheet_visibility depending on
    # workbook/version.
    #
    visibility = getattr(
        workbook,
        "sheet_visibility",
        None,
    )

    if not isinstance(
        visibility,
        (
            list,
            tuple,
        ),
    ):
        return "unknown"

    try:
        value = int(
            visibility[
                sheet_index_zero_based
            ]
        )
    except Exception:
        return "unknown"

    return {
        0: "visible",
        1: "hidden",
        2: "very_hidden",
    }.get(
        value,
        "unknown",
    )


def _iter_openpyxl_workbook(
    workbook_path: Path,
):
    try:
        from openpyxl import (
            load_workbook,
        )
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for "
            "XLSX/XLSM/XLTX/XLTM expansion."
        ) from exc

    #
    # data_only=True deliberately exposes cached/display
    # values rather than formula source text for Detection.
    #
    workbook = load_workbook(
        filename=str(
            workbook_path
        ),
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    try:
        for index, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            yield {
                "sheet_index": index,
                "sheet_name": (
                    worksheet.title
                ),
                "sheet_visibility": (
                    _openpyxl_visibility(
                        worksheet
                    )
                ),
                "rows": (
                    worksheet.iter_rows(
                        values_only=True
                    )
                ),
            }

    finally:
        workbook.close()


def _iter_xls_workbook(
    workbook_path: Path,
):
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError(
            "xlrd is required for "
            "legacy XLS expansion."
        ) from exc

    workbook = xlrd.open_workbook(
        str(
            workbook_path
        ),
        on_demand=True,
    )

    try:
        for sheet_index_zero in range(
            workbook.nsheets
        ):
            worksheet = (
                workbook.sheet_by_index(
                    sheet_index_zero
                )
            )

            def row_iterator(
                sheet=worksheet,
            ):
                for row_index in range(
                    sheet.nrows
                ):
                    yield (
                        sheet.row_values(
                            row_index
                        )
                    )

            yield {
                "sheet_index": (
                    sheet_index_zero
                    + 1
                ),
                "sheet_name": (
                    worksheet.name
                ),
                "sheet_visibility": (
                    _xls_visibility(
                        workbook,
                        sheet_index_zero,
                    )
                ),
                "rows": row_iterator(),
            }

            try:
                workbook.unload_sheet(
                    sheet_index_zero
                )
            except Exception:
                pass

    finally:
        try:
            workbook.release_resources()
        except Exception:
            pass


def _iter_xlsb_workbook(
    workbook_path: Path,
):
    try:
        from pyxlsb import (
            open_workbook,
        )
    except ImportError as exc:
        raise RuntimeError(
            "pyxlsb is required for "
            "XLSB expansion."
        ) from exc

    with open_workbook(
        str(
            workbook_path
        )
    ) as workbook:
        sheet_names = list(
            workbook.sheets
        )

        for index, sheet_name in enumerate(
            sheet_names,
            start=1,
        ):
            with workbook.get_sheet(
                sheet_name
            ) as worksheet:

                def row_iterator(
                    sheet=worksheet,
                ):
                    for row in sheet.rows():
                        yield [
                            getattr(
                                cell,
                                "v",
                                None,
                            )
                            for cell in row
                        ]

                yield {
                    "sheet_index": index,
                    "sheet_name": (
                        sheet_name
                    ),
                    #
                    # pyxlsb does not reliably expose workbook
                    # visibility state through its public API.
                    #
                    "sheet_visibility": (
                        "unknown"
                    ),
                    "rows": row_iterator(),
                }


def _iter_workbook_sheets(
    workbook_path: Path,
    extension: str,
):
    ext = _normalize_extension(
        extension
    )

    if ext in OPENPYXL_EXTENSIONS:
        yield from (
            _iter_openpyxl_workbook(
                workbook_path
            )
        )
        return

    if ext in XLRD_EXTENSIONS:
        yield from (
            _iter_xls_workbook(
                workbook_path
            )
        )
        return

    if ext in PYXLSB_EXTENSIONS:
        yield from (
            _iter_xlsb_workbook(
                workbook_path
            )
        )
        return

    raise RuntimeError(
        "Unsupported workbook extension "
        f"for expansion: {ext}"
    )
